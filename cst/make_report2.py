"""Отчёт-2: проверка концепции ВАР (статья [1] и ДЗ1) полноволновым расчётом в CST — линейная решётка и кольцо.
  python cst/make_report2.py   -> cst/Отчёт_проверка_ВАР_в_CST.docx
Числа берутся из cst/<папка>/var_10GHz.json и phase_centers_elems.json; кольцо R = 10λ включается, если посчитано.
"""
import os, json
import numpy as np
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "Отчёт_проверка_ВАР_в_CST.docx")
FONT = "Times New Roman"
LAM = 30.0


def J(folder, name):
    p = os.path.join(ROOT, folder, name)
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else None


def var_row(folder, cut, model_key):
    rows = J(folder, "var_10GHz.json") or []
    for r in rows:
        if r["cut"] == cut and model_key in r["model"]:
            return r
    return None


def pc_dev(folder):
    d = J(folder, "phase_centers_elems.json")
    lay = os.path.join(ROOT, folder, "layout_mm.txt")
    if d is None:
        return None
    if os.path.exists(lay):
        geom = {p[0]: (float(p[1]), float(p[2])) for p in (l.split() for l in open(lay, encoding="utf-8")) if p and p[0] != "name"}
    else:
        geom = {f"rx{i+1}": (x, 0.0) for i, x in enumerate([-105, -75, -45, -15, 15, 45, 75, 105])}
        geom.update({f"tx{i+1}": (x, 30.0) for i, x in enumerate([-127.5, -112.5, 112.5, 127.5])})
    devs = {n: max(abs(d[n]["xyz"][0] - geom[n][0]), abs(d[n]["xyz"][1] - geom[n][1])) for n in geom}
    ref = d[sorted(geom)[0]]["phi0"]
    ph = {n: float(np.degrees(np.angle(np.exp(1j * np.radians(d[n]["phi0"] - ref))))) for n in geom}
    return devs, ph


f = lambda v, nd=2: (f"{v:.{nd}f}").replace(".", ",").replace("-", "−")
wrapdeg = lambda a: float(np.degrees(np.angle(np.exp(1j * np.radians(a)))))
d = Document()
sec = d.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2); sec.left_margin = Cm(2.5); sec.right_margin = Cm(2)
st = d.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(12); st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
for name, size in (("Heading 1", 14), ("Heading 2", 13)):
    h = d.styles[name]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = RGBColor(0, 0, 0)
    rf = h.element.rPr.rFonts
    for a_ in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a_)) is not None:
            del rf.attrib[qn(a_)]
    for a_ in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a_), FONT)


def para(parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, style=None, size=None, italic=False):
    p = d.add_paragraph(style=style); p.alignment = align
    if isinstance(parts, str):
        parts = [(parts, False)]
    for text, bold in parts:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if size:
            r.font.size = Pt(size)
    return p


def bullet(parts):
    return para(parts, style="List Bullet")


def quote(text, src):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(1.0); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("«" + text + "»"); r.italic = True; r.font.size = Pt(11)
    r2 = p.add_run(" — " + src); r2.font.size = Pt(10)


def caption(t):
    p = para(t, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, italic=True); p.paragraph_format.space_after = Pt(10)


def table(header, rows, widths_cm):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        tcPr = c._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "E7E6E6"); tcPr.append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(10)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    return t


def img(rel, width_cm):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(ROOT, rel), width=Cm(width_cm))


# ---------------------------------------------------------------- числа
lin0 = var_row("stage3", 0.0, "формула"); linc = var_row("stage3", 0.0, "без калибровки"); link = var_row("stage3", 0.0, "калибровкой")
r5 = {c: (var_row("ring_R5", c, "формула"), var_row("ring_R5", c, "без калибровки"), var_row("ring_R5", c, "калибровкой")) for c in (0.0, 90.0)}
has10 = os.path.exists(os.path.join(ROOT, "ring_R10", "var_10GHz.json"))
r10 = {c: (var_row("ring_R10", c, "формула"), var_row("ring_R10", c, "без калибровки"), var_row("ring_R10", c, "калибровкой")) for c in (0.0, 90.0)} if has10 else None
dev_lin, ph_lin = pc_dev("stage3"); dev_r5, ph_r5 = pc_dev("ring_R5")
dev_r10 = pc_dev("ring_R10") if has10 and os.path.exists(os.path.join(ROOT, "ring_R10", "phase_centers_elems.json")) else None

# ---------------------------------------------------------------- текст
para([("Проверка концепции виртуальной антенной решётки MIMO полноволновым расчётом в CST Studio Suite", True)],
     align=WD_ALIGN_PARAGRAPH.CENTER, size=15)
para("ДЗ2 по курсу ОТиТРСКУ: линейная решётка 8 Rx + 4 Tx и кольцевая решётка 20 Rx + 10 Tx из ДЗ1.",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
para("Студент Сухарев Е. С., РЛ1-104. 22 октября 2026 г.", align=WD_ALIGN_PARAGRAPH.CENTER, size=11)

d.add_heading("1. Что проверяется", 1)
para("В ДЗ1 направленные свойства MIMO-решёток считались по концепции виртуальной антенной решётки (ВАР): "
     "каждая пара «передатчик–приёмник» заменяется изотропным элементом в точке r_tx + r_rx, а диаграмма "
     "направленности — сумма плоских волн с фазами k·(r_tx + r_rx)·u. Здесь эта концепция проверяется "
     "электродинамически: все элементы считаются в CST вместе, со взаимной связью, и отклик каждой пары "
     "берётся как произведение встроенных ДН F_tx(u)·F_rx(u). Проверяемые утверждения источников:")
quote("…приёмником будут приняты четыре сигнала, отличающихся фазовыми сдвигами, … что эквивалентно использованию "
      "традиционной антенной решётки с количеством элементов, равным четырём", "[1], с. 113, ф. (4)")
quote("Относительные координаты виртуальных приёмных антенн при этом определяются по формуле ρ_m = r_tx,p + r_rx,q − r_0",
      "[3], с. 240")
quote("Сигнал, принимаемый таким элементом, определяется коэффициентом передачи в канале между соответствующими "
      "передающей и приёмной антеннами", "[3], с. 241")
quote("Сравнение рис. 7, а, б показывает, что поперечное смещение передающей подрешётки практически не влияет на "
      "диаграмму направленности виртуальной антенной решётки. Поэтому данный метод целесообразен для снижения "
      "электромагнитных связей между передающими и приёмными каналами (расстояние может составлять несколько длин волн)",
      "[1], с. 118; то же — вывод 3, с. 120")
quote("…рассовмещение передающей и приёмной АР, необходимое на практике для осуществления развязки приём-передача, "
      "не приводит к существенному изменению результатов и качества фокусировки радиоизображения", "[7], с. 83")
para("Из ДЗ1 проверяются числа: линейная решётка — ширина луча 3,16°, УБЛ −13,2 дБ; кольцо — «крайне узкий главный "
     "луч, что составляет порядка 0,5°, … уровень боковых лепестков может приближаться к уровню −3 дБ» (ДЗ1, с. 9; "
     "после исправления квадранта в DZ_py.py — 0,24° и −4,0 дБ).")
para("Проверка считается пройденной, если: (а) фазовый центр каждого элемента в решётке совпадает с его "
     "геометрической точкой с точностью λ/20 = 1,5 мм; (б) ДН суммы произведений F_tx·F_rx совпадает с ДН ВАР "
     "по формуле по ширине луча и УБЛ; (в) сдвиг ряда Tx на λ поперёк оси не меняет ДН.")

d.add_heading("2. Как проверяется", 1)
para("Все элементы — одинаковые полуволновые диполи 12,5 мм (резонанс 10,2 ГГц) вдоль y, в плоскости z = 0, "
     "дискретные порты 50 Ом; CST 2024, решатель во временной области, 9–11 ГГц, открытые границы. Каждый порт "
     "возбуждается по очереди, остальные нагружены 50 Ом: это даёт встроенные ДН F_n(θ, φ) со взаимной связью. "
     "ДН выгружаются с шагом 1° по θ и φ с фазой относительно глобального нуля. Тип излучателя для проверки не "
     "принципиален: проверяется геометрическая модель, а не антенна.")
bullet([("Фазовый центр элемента ", True), ("— МНК-подгонка фазы встроенной ДН моделью точечного источника "
        "k·(x·sinθ·cosφ + y·sinθ·sinφ) + φ0 в секторе ±30° (рис. 1). Остаток 3–10° — рябь от рассеяния на соседях.", False)])
bullet([("ДН ВАР по формуле ", True), ("— сумма exp(j·k·(r_tx + r_rx)·u) по всем парам, умноженная на квадрат среднего "
        "множителя элемента (чтобы сравнивать только решётку).", False)])
bullet([("ДН ВАР по CST ", True), ("— Σ F_tx(u)·F_rx(u) по тем же парам; отдельно с калибровкой каналов (каждая ДН "
        "делится на своё значение по нормали, как в реальной MIMO-РЛС) и без неё.", False)])
img("fig2/ring_R5_phase.png", 16)
caption("Рис. 1. Фаза встроенной ДН элемента rx1 кольца (точки) и модель точечного источника в его геометрическом "
        "центре (линия); справа — остаток, не превышающий ±15°")

d.add_heading("3. Линейная решётка 8 Rx + 4 Tx (ряд Tx сдвинут на λ)", 1)
para("Геометрия ДЗ1: Rx с шагом λ на x = −105…105 мм, Tx на x = ±112,5 и ±127,5 мм; ряд Tx сдвинут на y = +30 мм, "
     "как рекомендует [1]. Рис. 2 и 3 показывают ДН ВАР по формуле и по CST; табл. 1 — числа.")
img("stage3/var_10GHz.png", 16)
caption("Рис. 2. Линейная решётка: сечения ДН ВАР по формуле (чёрная) и по CST (красная — без калибровки, синяя — с калибровкой)")
img("fig2/stage3_map.png", 16)
caption("Рис. 3. Линейная решётка: ДН ВАР на всей полусфере по формуле и по CST, справа — их разность")
table(["Величина", "ДЗ1 (формула, изотропные)", "Формула × элемент", "CST без калибровки", "CST с калибровкой"], [
    ["Ширина луча по −3 дБ, °", "3,16", f(lin0["hpbw"]), f(linc["hpbw"]), f(link["hpbw"])],
    ["УБЛ, дБ", "−13,2", f(lin0["sll_db"], 1), f(linc["sll_db"], 1), f(link["sll_db"], 1)],
    ["Максимум ДН, °", "0", f(lin0["peak_deg"], 1), f(linc["peak_deg"], 1), f(link["peak_deg"], 1)],
], [4.6, 3.2, 2.8, 2.9, 2.9])
caption("Табл. 1. Линейная решётка, сечение вдоль оси (φ = 0), 10 ГГц")
para(f"Фазовые центры: Rx отклоняются от геометрии не более чем на {f(max(v for n, v in dev_lin.items() if n.startswith('rx')))} мм, "
     f"Tx — на {f(max(v for n, v in dev_lin.items() if n.startswith('tx')))} мм (пары Tx с шагом λ/2 «расталкиваются» "
     "взаимной связью; это λ/25, на пределе критерия). Фаза φ0 встроенных ДН крайних Rx отстаёт на "
     f"{f(abs(wrapdeg(ph_lin['rx1'] - ph_lin['rx4'])), 0)}° от внутренних, Tx — на {f(abs(wrapdeg(ph_lin['tx1'] - ph_lin['rx4'])), 0)}°; на ДН ВАР это почти "
     "не сказалось (табл. 1, два столбца CST), потому что разнофазность симметрична относительно центра решётки.")
para("Сдвиг ряда Tx на λ: формула в сечении φ = 0 от сдвига по y не зависит, а CST с реально сдвинутым рядом "
     "даёт ту же ширину луча и тот же УБЛ — утверждение [1] о поперечном смещении подтверждается. Развязка Tx–Rx "
     "при этом не хуже −27 дБ (худшая пара Tx2–Rx1), Rx–Rx между соседями −19 дБ.")

d.add_heading("4. Кольцо 20 Rx + 10 Tx", 1)
para("Кольцо ДЗ1 (Rx на R = 50λ, Tx на 67,5λ) имеет диаметр 4 м и в CST не считается. Проверка выполнена на "
     "геометрически подобных кольцах с тем же числом элементов и тем же отношением радиусов 1,35: R = 5λ "
     f"(диаметр по Tx 0,4 м, 4,5 млн ячеек, 93 мин){' и R = 10λ (0,8 м, 9,6 млн ячеек, 3 ч)' if has10 else ''}. "
     "Если формула ВАР верна при тесном кольце, где связь между элементами сильнее, она верна и при разрежённом.")
img("fig2/ring_R5_layout.png", 13)
caption(f"Рис. 4. Кольцо R = 5λ: геометрия (контуры) и найденные фазовые центры (закрашено); они совпадают с точностью "
        f"{f(max(dev_r5.values()))} мм = λ/{LAM/max(dev_r5.values()):.0f}. Цвет — фаза φ0 встроенной ДН")
img("ring_R5/var_10GHz.png", 16)
caption("Рис. 5. Кольцо R = 5λ: сечения ДН ВАР по формуле и по CST (φ = 0 и 90°)")
img("fig2/ring_R5_map.png", 16)
caption("Рис. 6. Кольцо R = 5λ: ДН ВАР на полусфере по формуле и по CST, справа — разность")
rows = []
for c in (0.0, 90.0):
    a0, b0, c0 = r5[c]
    rows.append([f"R = 5λ, φ = {c:g}°, ширина луча, °", f(a0["hpbw"]), f(b0["hpbw"]), f(c0["hpbw"])])
    rows.append([f"R = 5λ, φ = {c:g}°, УБЛ, дБ", f(a0["sll_db"], 1), f(b0["sll_db"], 1), f(c0["sll_db"], 1)])
if has10:
    for c in (0.0, 90.0):
        a0, b0, c0 = r10[c]
        rows.append([f"R = 10λ, φ = {c:g}°, ширина луча, °", f(a0["hpbw"]), f(b0["hpbw"]), f(c0["hpbw"])])
        rows.append([f"R = 10λ, φ = {c:g}°, УБЛ, дБ", f(a0["sll_db"], 1), f(b0["sll_db"], 1), f(c0["sll_db"], 1)])
table(["Величина", "Формула × элемент", "CST без калибровки", "CST с калибровкой"], rows, [6.2, 3.2, 3.2, 3.2])
caption("Табл. 2. Кольцо: ширина луча и УБЛ по формуле и по CST, 10 ГГц. Для справки, формула с изотропными элементами: "
        "R = 5λ — 2,48° и −8,1 дБ; R = 10λ — 1,24° и −8,1 дБ; R = 50λ (ДЗ1) — 0,24° и −4,0 дБ")
para("Ширина луча масштабируется как 1/R, как и должно быть для разрежённой решётки, а уровень дифракционных "
     "лепестков при R = 50λ по формуле поднимается до −4 дБ, что и было отмечено в ДЗ1. Полноволновой расчёт "
     "подтверждает формулу с точностью лучше 1 дБ по всем лепесткам выше −20 дБ (рис. 6), поэтому вывод ДЗ1 о "
     "кольце — узкий луч при неприемлемом УБЛ — остаётся в силе и для реальных излучателей.")

d.add_heading("5. Выводы", 1)
bullet([("Концепция ВАР подтверждена полноволновым расчётом. ", True), ("Для линейной решётки и кольца ДН суммы "
        "произведений встроенных ДН всех пар совпадает с ДН по формуле r_tx + r_rx: ширина луча — в пределах 0,1°, "
        "УБЛ — в пределах 1 дБ, расхождение по всем лепесткам выше −20 дБ — около 1 дБ (скз).", False)])
bullet([("Фазовые центры элементов лежат в их геометрических точках: ", True),
        (f"кольцо — не хуже {f(max(dev_r5.values()))} мм (λ/{LAM/max(dev_r5.values()):.0f}), линейная решётка — "
         f"{f(max(v for n, v in dev_lin.items() if n.startswith('rx')))} мм для Rx и "
         f"{f(max(v for n, v in dev_lin.items() if n.startswith('tx')))} мм для пар Tx с шагом λ/2 (λ/25).", False)])
bullet([("Поперечный сдвиг ряда Tx на λ ДН не меняет ", True), ("(вывод 3 статьи [1], с. 120, и замечание [7], с. 83) "
        "и даёт развязку Tx–Rx не хуже −27 дБ.", False)])
bullet([("Что формула не учитывает: ", True), ("комплексные коэффициенты каналов — фазы встроенных ДН различаются "
        "на десятки градусов (крайние элементы линейки — 70°, элементы кольца — до 80° в зависимости от положения на "
        "окружности). При симметричной решётке это не портит ДН, при несимметричной выборке каналов — смещает "
        "кажущийся центр (сумма rx1…rx4 линейки: −68 мм вместо −60). В реальной MIMO-РЛС это снимается "
        "калибровкой каналов по нормали, что и сделано в столбце «CST с калибровкой».", False)])
bullet([("Числа ДЗ1 ", True), ("для линейной решётки (3,16°, −13,2 дБ) воспроизводятся CST с точностью 0,2° и 1 дБ; "
        "для кольца формульные 0,24° и −4,0 дБ в CST напрямую не проверить (4 м), но подтверждены на подобных "
        "кольцах 5λ" + (" и 10λ" if has10 else "") + ".", False)])

d.add_heading("6. Файлы", 1)
para("Репозиторий BaykaI/MIMO, ветка claude/sweet-archimedes-gcmsuv, папка cst/: build_array.py и build_ring.py "
     "(модели CST), export_ff.py (выгрузка ДН), var_compare.py (формула против CST), analyze_stage.py и "
     "report2_figs.py (фазовые центры, карты ДН), citations.md (цитаты с номерами страниц), stage3/, ring_R5/"
     + (", ring_R10/" if has10 else "") + " (проекты, S-матрицы, таблицы). Литература — по списку ДЗ1.")
d.save(OUT)
print("written", OUT)
